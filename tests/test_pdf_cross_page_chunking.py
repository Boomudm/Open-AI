"""Phase 4 — cross-page PDF chunking and parsing quality."""

from __future__ import annotations

from pathlib import Path

import pytest

from spark.ingestion.chunk.registry import chunk_text_leaves
from spark.ingestion.chunking.context_header import build_chunk_header, embed_text_with_header
from spark.ingestion.chunking.page_markers import concat_pages_with_markers, page_for_char_offset
from spark.ingestion.chunking.structure import split_structure_aware
from spark.ingestion.ocr_quality import OCRQualityTracker, chunk_index_rejected
from spark.ingestion.parsers.word_parser import WordParser, _extract_structured_blocks
from spark.ingestion.validators import Validator


FIXTURES = Path(__file__).resolve().parent / "fixtures" / "parsing"


def test_cross_page_procedure_not_split_mid_step() -> None:
    raw = (FIXTURES / "multi_page_procedure.txt").read_text(encoding="utf-8")
    pages = [
        (1, raw.split("<!-- page 2 -->")[0].replace("<!-- page 1 -->", "").strip()),
        (2, raw.split("<!-- page 2 -->")[1].strip()),
    ]
    full_text, markers = concat_pages_with_markers(pages)
    leaves = chunk_text_leaves(
        full_text,
        strategy="structure_aware",
        chunk_size=400,
        chunk_overlap=40,
    )
    texts = [leaf.text for leaf in leaves if leaf.text.strip()]
    assert texts
    joined = "\n".join(texts)
    assert "Inspect the cooling fan" in joined
    assert "Shut down the machine" in joined
    step_three_offset = full_text.find("3. Inspect the cooling fan")
    assert page_for_char_offset(markers, step_three_offset) == 2


def test_structure_aware_splits_procedure_steps() -> None:
    text = (FIXTURES / "word_sop_procedure.txt").read_text(encoding="utf-8")
    chunks = split_structure_aware(text, chunk_size=500, chunk_overlap=50)
    assert len(chunks) >= 3
    assert any("grease fitting" in c for c in chunks)


def test_context_header_stable_and_embed_enriched() -> None:
    ctx = {
        "machine_id": "MC-001",
        "document_type": "machine_manual",
        "title": "Servo Manual",
    }
    payload = {"page": 13, "section_path": "Cooling > Fan", "chunk_level": "child"}
    header = build_chunk_header(ctx, payload)
    assert "MC-001" in header
    assert "machine_manual" in header
    assert "Page 13" in header
    assert "Cooling > Fan" in header
    assert "2026" not in header
    enriched = embed_text_with_header("Check bearing torque.", ctx, payload)
    assert enriched.startswith(header)
    assert "Check bearing torque." in enriched


def test_ocr_index_rejection() -> None:
    assert chunk_index_rejected(65.0, "ocr") is True
    assert chunk_index_rejected(75.0, "ocr") is False
    assert chunk_index_rejected(65.0, "pypdf") is False


def test_ocr_job_quality_insufficient_for_manual_returns_tag_not_raise() -> None:
    tracker = OCRQualityTracker()
    for page in range(1, 5):
        tracker.add(page, "ocr", 40.0)
    detail = tracker.enforce_job_quality(document_type="machine_manual", total_page_count=4)
    assert detail is not None


@pytest.mark.asyncio
async def test_validator_rejects_ocr_low_confidence_chunk() -> None:
    from spark.ingestion.parsers.base import ParsedRecord

    validator = Validator()
    rec = ParsedRecord(
        route="qdrant",
        payload={
            "page_content": "low quality",
            "ocr_rejected": True,
            "page_ocr_confidence": 55.0,
        },
        text_for_embedding="low quality",
    )
    ok, code, _ = validator.validate(rec, {}, {})
    assert ok is False
    assert code == "ocr_low_confidence"


@pytest.mark.asyncio
async def test_word_parser_parent_child(tmp_path: Path) -> None:
    try:
        from docx import Document as DocxDocument
    except ImportError:
        pytest.skip("python-docx not installed")

    path = tmp_path / "sop.docx"
    doc = DocxDocument()
    long_body = "\n\n".join(f"Step {i}: " + ("detail " * 80) for i in range(1, 8))
    doc.add_paragraph(long_body)
    doc.save(str(path))

    parser = WordParser()
    config = {
        "chunking": {
            "mode": "parent_child",
            "strategy": "parent_child",
            "parent_chunk_size": 600,
            "child_chunk_size": 120,
            "child_chunk_overlap": 20,
        }
    }
    records = [rec async for rec in parser.parse(path, config)]
    assert records
    assert all(rec.payload.get("parent_chunk_id") for rec in records)
    assert all(rec.payload.get("parent_content") for rec in records)


def test_word_doc_extension_rejected(tmp_path: Path) -> None:
    from spark.ingestion.parsers.base import ParserError

    legacy = tmp_path / "legacy.doc"
    legacy.write_bytes(b"fake")
    parser = WordParser()

    async def _run():
        async for _ in parser.parse(legacy, {}):
            pass

    with pytest.raises(ParserError) as exc:
        import asyncio

        asyncio.run(_run())
    assert exc.value.reason_code == "unsupported_file_type"


def test_word_table_extraction() -> None:
    try:
        from docx import Document as DocxDocument
    except ImportError:
        pytest.skip("python-docx not installed")
    import tempfile

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        path = Path(tmp.name)
    doc = DocxDocument()
    doc.add_paragraph("Intro paragraph")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Part"
    table.cell(0, 1).text = "Torque"
    table.cell(1, 0).text = "M6"
    table.cell(1, 1).text = "12 Nm"
    doc.save(str(path))
    blocks, tables = _extract_structured_blocks(path)
    body = "\n".join(b.text for b in blocks if b.text)
    assert "Intro paragraph" in body
    assert tables and "M6" in tables[0].markdown
    path.unlink(missing_ok=True)
